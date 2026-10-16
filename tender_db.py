"""
標案資料庫：只保存「與本司業務相關」（符合關鍵字規則）的案件

  data/db/YYYY-MM.csv      標案資料庫（只含符合案件），依公告月份分檔
  data/raw/YYYY-MM.csv     原始收錄（全部勞務類公告），僅供日後調整關鍵字時重新篩選，不對外顯示
  data/標案資料庫.xlsx      每次執行重新產生：標案資料庫 / 月份統計 / 說明
"""
from __future__ import annotations

import csv
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent
DB_DIR = ROOT / "data" / "db"
RAW_DIR = ROOT / "data" / "raw"
XLSX = ROOT / "data" / "標案資料庫.xlsx"

FIELDS = ["主鍵", "公告日", "機關", "案號", "案名", "公告類型", "傳輸次數", "採購性質",
          "預算", "截止投標", "符合", "關鍵字群組", "命中關鍵字", "分數", "連結",
          "首次收錄", "最後更新", "資料來源"]
MATCH_COLS = ("符合", "關鍵字群組", "命中關鍵字", "分數")


def row_key(r: dict) -> str:
    return f"{r['unit_name']}|{r['job_number']}|{r.get('seq') or '01'}|{r['type']}"


def _roc_to_iso(s: str) -> str:
    """'115/10/20 17:00' → '2026-10-20 17:00'（方便排序，也避免試算表誤判）"""
    import re
    m = re.match(r"\s*(\d{2,3})/(\d{1,2})/(\d{1,2})(.*)", s or "")
    if not m:
        return s or ""
    return f"{int(m[1]) + 1911}-{int(m[2]):02d}-{int(m[3]):02d}{m[4].rstrip()}"


def _match_cols(m: dict | None) -> dict:
    return {
        "符合": "Y" if m else "",
        "關鍵字群組": "/".join(m["groups"]) if m else "",
        "命中關鍵字": "、".join(m["keywords"]) if m else "",
        "分數": m["score"] if m else "",
    }


def _to_row(r: dict, m: dict | None, stamp: str, url: str, budget) -> dict:
    d = r.get("date") or ""
    return {
        "主鍵": row_key(r),
        "公告日": f"{d[:4]}-{d[4:6]}-{d[6:]}" if len(d) == 8 else d,
        "機關": r["unit_name"],
        "案號": r["job_number"],
        "案名": r["title"],
        "公告類型": r["type"],
        "傳輸次數": r.get("seq", ""),
        "採購性質": r.get("category", ""),
        "預算": "" if budget is None else budget,
        "截止投標": _roc_to_iso(r.get("deadline_list", "")),
        **_match_cols(m),
        "連結": url,
        "首次收錄": stamp,
        "最後更新": stamp,
        "資料來源": r.get("source", "today"),
    }


def _load(path: Path) -> dict[str, dict]:
    if not path.exists():
        return {}
    with path.open(encoding="utf-8-sig", newline="") as f:
        rows = {row["主鍵"]: row for row in csv.DictReader(f)}
    for row in rows.values():  # 舊資料民國日期 → 西元
        row["截止投標"] = _roc_to_iso(row.get("截止投標", ""))
    return rows


def _save(path: Path, rows: dict[str, dict]) -> None:
    if not rows:
        if path.exists():
            path.unlink()
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    ordered = sorted(rows.values(), key=lambda x: (x["公告日"], x["機關"], x["案號"]), reverse=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS, extrasaction="ignore")
        w.writeheader()
        w.writerows(ordered)


def _month(row: dict) -> str:
    return (row.get("公告日") or "")[:7] or "unknown"


def _sync_db(raw_rows: list[dict]) -> tuple[int, int, int, list[dict]]:
    """依原始資料的「符合」欄更新標案資料庫：符合→新增/更新，不符合→移除。
    回傳 (新增, 更新, 移除, 需同步到雲端的列)。"""
    by_month: dict[str, list[dict]] = defaultdict(list)
    for r in raw_rows:
        by_month[_month(r)].append(r)
    added = updated = removed = 0
    to_sheet: list[dict] = []
    for month, items in by_month.items():
        path = DB_DIR / f"{month}.csv"
        db = _load(path)
        for r in items:
            k = r["主鍵"]
            if r.get("符合") == "Y":
                if k not in db:
                    added += 1
                    to_sheet.append(r)
                elif any(str(r.get(c, "")) != str(db[k].get(c, "")) for c in FIELDS if c != "最後更新"):
                    updated += 1
                    to_sheet.append(r)
                db[k] = r
            elif k in db:
                del db[k]
                removed += 1
                to_sheet.append(r)   # 符合欄為空 → 雲端端會移除（未填備註者）
        _save(path, db)
    return added, updated, removed, to_sheet


def upsert(recs: list[dict], match_fn, url_fn, budget_fn, now: datetime):
    """寫入原始收錄與標案資料庫。回傳 (新增, 更新, 需同步到雲端的列, 移除)。"""
    stamp = now.strftime("%Y-%m-%d %H:%M")
    by_month: dict[str, list[dict]] = defaultdict(list)
    for r in recs:
        d = r.get("date") or now.strftime("%Y%m%d")
        by_month[f"{d[:4]}-{d[4:6]}"].append(r)

    changed_raw: list[dict] = []
    for month, items in by_month.items():
        path = RAW_DIR / f"{month}.csv"
        raw = _load(path)
        for r in items:
            m = match_fn(r)
            new = _to_row(r, m, stamp, url_fn(r, m), budget_fn(r))
            k = new["主鍵"]
            if k in raw:
                old = raw[k]
                new["首次收錄"] = old.get("首次收錄") or stamp
                # 備援來源（今日招標頁）沒有預算，不覆蓋查詢頁已有的資料
                for col in ("預算", "採購性質", "傳輸次數"):
                    if not new[col] and old.get(col):
                        new[col] = old[col]
                if any(str(new[c]) != str(old.get(c, "")) for c in FIELDS if c not in ("最後更新", "首次收錄")):
                    raw[k] = new
                    changed_raw.append(new)
            else:
                raw[k] = new
                changed_raw.append(new)
        _save(path, raw)
    added, updated, removed, to_sheet = _sync_db(changed_raw)
    return added, updated, to_sheet, removed


def rescore(match_fn, now: datetime) -> tuple[int, list[dict], tuple[int, int, int]]:
    """依目前關鍵字規則重新篩選全部原始收錄，重建標案資料庫。
    回傳 (原始總筆數, 需同步到雲端的列, (新增, 更新, 移除))。"""
    stamp = now.strftime("%Y-%m-%d %H:%M")
    total, all_raw = 0, []
    for path in sorted(RAW_DIR.glob("*.csv")):
        raw = _load(path)
        dirty = False
        for row in raw.values():
            total += 1
            m = match_fn({"type": row["公告類型"], "title": row["案名"], "unit_name": row["機關"], "pk": ""})
            new = _match_cols(m)
            if any(str(new[c]) != str(row.get(c, "")) for c in MATCH_COLS):
                row.update(new)
                row["最後更新"] = stamp
                dirty = True
            all_raw.append(row)
        if dirty:
            _save(path, raw)
    # 資料庫裡有、但原始收錄已不符合者也要移除（_sync_db 會處理）
    added, updated, removed, to_sheet = _sync_db(all_raw)
    return total, to_sheet, (added, updated, removed)


def load_all() -> list[dict]:
    rows = []
    for p in sorted(DB_DIR.glob("*.csv")):
        rows.extend(_load(p).values())
    return rows


def raw_counts() -> Counter:
    c = Counter()
    for p in sorted(RAW_DIR.glob("*.csv")):
        for row in _load(p).values():
            c[_month(row)] += 1
    return c


def build_xlsx(now: datetime, recent_days: int = 30) -> Path | None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("ℹ️  未安裝 openpyxl，略過 Excel")
        return None

    def _score(x):
        try:
            return int(float(x.get("分數") or 0))
        except ValueError:
            return 0

    rows = load_all()
    rows.sort(key=lambda x: (x["公告日"], _score(x)), reverse=True)

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1D4ED8")
    head_font = Font(color="FFFFFF", bold=True)
    cols = [("公告日", 11), ("機關", 26), ("案名", 60), ("預算", 14), ("截止投標", 16),
            ("公告類型", 24), ("傳輸次數", 8), ("案號", 18), ("關鍵字群組", 18),
            ("命中關鍵字", 24), ("分數", 6), ("連結", 10), ("首次收錄", 16)]

    ws = wb.active
    ws.title = "標案資料庫"
    ws.append([c for c, _ in cols])
    for i, (c, w) in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i)
        cell.fill, cell.font = head_fill, head_font
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    link_ci = [c for c, _ in cols].index("連結") + 1
    for r in rows:
        vals = []
        for c, _ in cols:
            v = r.get(c, "")
            if c in ("預算", "分數") and str(v).strip():
                try:
                    v = int(float(v))
                except ValueError:
                    pass
            vals.append("開啟" if c == "連結" and v else v)
        ws.append(vals)
        n = ws.max_row
        if r.get("連結"):
            lc = ws.cell(row=n, column=link_ci)
            lc.hyperlink, lc.font = r["連結"], Font(color="1D4ED8", underline="single")
        ws.cell(row=n, column=4).number_format = "#,##0"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    ws = wb.create_sheet("月份統計")
    ws.append(["月份", "掃描公告數", "符合件數", "符合案預算合計(元)"])
    for i in range(1, 5):
        ws.cell(row=1, column=i).fill, ws.cell(row=1, column=i).font = head_fill, head_font
        ws.column_dimensions[get_column_letter(i)].width = 20
    scanned = raw_counts()
    hit_n, hit_b = Counter(), Counter()
    for r in rows:
        mth = _month(r)
        hit_n[mth] += 1
        try:
            hit_b[mth] += int(float(r.get("預算") or 0))
        except ValueError:
            pass
    for mth in sorted(set(scanned) | set(hit_n), reverse=True):
        ws.append([mth, scanned[mth], hit_n[mth], hit_b[mth]])
        ws.cell(row=ws.max_row, column=4).number_format = "#,##0"

    ws = wb.create_sheet("說明")
    for line in [
        "政府標案資料庫（自動產生，請勿手動編輯此檔；每次執行會覆蓋）",
        f"產生時間：{now:%Y-%m-%d %H:%M}　資料庫共 {len(rows)} 筆",
        "只收錄與本司業務相關的案件（案名符合 config.yml 關鍵字規則、未被排除詞排除、不含更正公告）",
        "資料來源：政府電子採購網「標案查詢」（勞務類）",
        "預算空白：機關未公開預算，或該筆來自「今日招標」備援頁",
    ]:
        ws.append([line])
    ws.column_dimensions["A"].width = 100

    XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(XLSX)
    return XLSX
